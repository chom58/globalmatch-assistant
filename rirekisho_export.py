"""履歴書（JIS形式）Excel 出力モジュール

英語CVから抽出・候補者が補完した構造化データを、
`templates/rirekisho_a4.xlsx`（A4 2枚・左右見開き）に流し込んで xlsx bytes を返す。
UI（Streamlit）には依存しない純関数のみ。PDF 変換は LibreOffice(soffice) がある環境でのみ動く。

データ仕様（build_rirekisho_xlsx の data 引数）:
    {
        "date": date | "YYYY-MM-DD" | None,     # 記入日（None なら今日）
        "name": str, "name_kana": str,
        "birth_date": date | "YYYY-MM-DD" | None,
        "gender": str | None,                    # 任意（空欄可）
        "postal_code": str, "address": str, "address_kana": str,
        "phone": str, "email": str,
        "contact": {"postal_code","address","address_kana","phone","email"} | None,
        "education":   [{"year": int|None, "month": int|None, "text": str}],
        "work_history": [{"year": int|None, "month": int|None, "text": str}],
        "qualifications": [{"year": int|None, "month": int|None, "text": str}],
        "motivation": str, "wishes": str,
    }
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, Side

TEMPLATE_PATH = Path(__file__).parent / "templates" / "rirekisho_a4.xlsx"
SHEET_NAME = "履歴書"

# --- セル配置（テンプレ実測） ---------------------------------------------
CELLS = {
    "date": "E3",             # 「　年　月　日現在」
    "name_kana": "C6",
    "name": "C9",
    "birth": "B14",           # 「年　月　日生（満　歳）」
    "gender": "F14",
    "address_kana": "C16",
    "postal_code": "C19",     # 「〒」
    "address": "C21",
    "phone": "I16",
    "email": "H21",
    "contact_kana": "C25",
    "contact_postal": "C28",
    "contact_address": "C30",
    "contact_phone": "I25",
    "contact_email": "H30",
    "motivation": "L47",
}

# 学歴・職歴: 左ページ 16 行 → 右ページ上部 6 行（年, 月, 内容）
HISTORY_SLOTS = (
    [(f"B{r}", f"C{r}", f"D{r}") for r in range(38, 81, 3)]
    + [("B83", "C83", "D83")]
    + [(f"L{r}", f"M{r}", f"N{r}") for r in (5, 8, 11, 14, 16, 19)]
)
MAX_HISTORY_ROWS = len(HISTORY_SLOTS)  # 22

# 資格・免許: 右ページ 6 行
QUALIFICATION_SLOTS = [(f"L{r}", f"M{r}", f"N{r}") for r in (25, 28, 31, 34, 37, 40)]
MAX_QUALIFICATION_ROWS = len(QUALIFICATION_SLOTS)

# 本人希望記入欄: 5 ブロック（1 ブロック 1 行として流す）
WISH_SLOTS = ["L71", "L74", "L77", "L80", "L83"]

# 写真枠（テンプレの図形は openpyxl で保持できないため自前で描く）
PHOTO_ANCHOR = "H4"
PHOTO_RANGE = "H4:I14"
PHOTO_WIDTH_PX = 105   # 約 28mm
PHOTO_HEIGHT_PX = 140  # 約 37mm

PRINT_SCALE = 85  # テンプレ既定 90% だと LibreOffice で縦がはみ出す


# --- ユーティリティ ---------------------------------------------------------
def parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m", "%Y/%m"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def compute_age(birth: date, on: date) -> int:
    """満年齢（on 時点）"""
    age = on.year - birth.year
    if (on.month, on.day) < (birth.month, birth.day):
        age -= 1
    return age


def _fmt_ym(year, month) -> tuple[str, str]:
    y = "" if year in (None, "") else str(int(year))
    m = "" if month in (None, "") else str(int(month))
    return y, m


def build_history_rows(education: list[dict], work_history: list[dict]) -> tuple[list[dict], int]:
    """学歴・職歴を JIS 慣行（見出し→行→「以上」）で 1 本の行リストにする。

    Returns: (rows, overflow)  overflow は上限超過で落とした行数
    """
    rows: list[dict] = []
    if education:
        rows.append({"year": None, "month": None, "text": "学歴", "align": "center"})
        rows.extend({**e, "align": "left"} for e in education if (e.get("text") or "").strip())
    if work_history:
        rows.append({"year": None, "month": None, "text": "職歴", "align": "center"})
        rows.extend({**w, "align": "left"} for w in work_history if (w.get("text") or "").strip())
    if not rows:
        return [], 0
    closing = {"year": None, "month": None, "text": "以上", "align": "right"}
    overflow = 0
    if len(rows) + 1 > MAX_HISTORY_ROWS:
        overflow = len(rows) + 1 - MAX_HISTORY_ROWS
        rows = rows[: MAX_HISTORY_ROWS - 1]
    rows.append(closing)
    return rows, overflow


def _set_aligned(ws, coord: str, value, horizontal: str | None = None) -> None:
    cell = ws[coord]
    cell.value = value or None
    if horizontal:
        al = copy(cell.alignment)
        cell.alignment = Alignment(
            horizontal=horizontal,
            vertical=al.vertical or "center",
            wrap_text=al.wrap_text,
            indent=al.indent,
        )


def _prepare_photo(photo_bytes: bytes) -> io.BytesIO:
    """3:4 に中央クロップして縮小、PNG bytes を返す"""
    from PIL import Image  # python-pptx 経由で導入済み。ここで遅延 import

    img = Image.open(io.BytesIO(photo_bytes)).convert("RGB")
    target_ratio = PHOTO_WIDTH_PX / PHOTO_HEIGHT_PX
    w, h = img.size
    if w / h > target_ratio:  # 横長 → 左右をカット
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        img = img.crop((left, 0, left + new_w, h))
    else:  # 縦長 → 上下をカット
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        img = img.crop((0, top, w, top + new_h))
    img = img.resize((PHOTO_WIDTH_PX * 3, PHOTO_HEIGHT_PX * 3), Image.LANCZOS)  # 印刷向けに高解像度
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def _draw_photo_placeholder(ws) -> None:
    ws.merge_cells(PHOTO_RANGE)
    dashed = Side(style="dashed", color="808080")
    top_left = ws[PHOTO_ANCHOR]
    top_left.value = "写真をはる位置\n\n縦 36〜40mm\n横 24〜30mm\n本人単身胸から上"
    top_left.font = Font(name="ＭＳ Ｐ明朝", size=7, color="808080")
    top_left.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    min_col, min_row, max_col, max_row = 8, 4, 9, 14  # H4:I14
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(
                left=dashed if cell.column == min_col else None,
                right=dashed if cell.column == max_col else None,
                top=dashed if cell.row == min_row else None,
                bottom=dashed if cell.row == max_row else None,
            )


# --- メイン ---------------------------------------------------------------
def build_rirekisho_xlsx(data: dict, photo_bytes: bytes | None = None) -> tuple[bytes, dict]:
    """履歴書 xlsx を生成する。

    Returns: (xlsx_bytes, info)  info = {"history_overflow": int, "qualification_overflow": int}
    """
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME]

    fill_date = parse_date(data.get("date")) or date.today()
    ws[CELLS["date"]].value = f"{fill_date.year}年{fill_date.month:>3}月{fill_date.day:>3}日現在"

    ws[CELLS["name_kana"]].value = (data.get("name_kana") or "").strip() or None
    ws[CELLS["name"]].value = (data.get("name") or "").strip() or None

    birth = parse_date(data.get("birth_date"))
    if birth:
        age = compute_age(birth, fill_date)
        ws[CELLS["birth"]].value = (
            f"{birth.year}年{birth.month:>3}月{birth.day:>3}日生　（満{age:>3}歳）"
        )
    gender = (data.get("gender") or "").strip()
    ws[CELLS["gender"]].value = gender or "※性別"

    ws[CELLS["address_kana"]].value = (data.get("address_kana") or "").strip() or None
    postal = (data.get("postal_code") or "").strip()
    ws[CELLS["postal_code"]].value = f"〒 {postal}" if postal else "〒"
    ws[CELLS["address"]].value = (data.get("address") or "").strip() or None
    ws[CELLS["phone"]].value = (data.get("phone") or "").strip() or None
    ws[CELLS["email"]].value = (data.get("email") or "").strip() or None

    contact = data.get("contact") or {}
    if any((contact.get(k) or "").strip() for k in ("address", "phone", "email")):
        ws[CELLS["contact_kana"]].value = (contact.get("address_kana") or "").strip() or None
        c_postal = (contact.get("postal_code") or "").strip()
        ws[CELLS["contact_postal"]].value = f"〒 {c_postal}" if c_postal else "〒"
        ws[CELLS["contact_address"]].value = (contact.get("address") or "").strip() or None
        ws[CELLS["contact_phone"]].value = (contact.get("phone") or "").strip() or None
        ws[CELLS["contact_email"]].value = (contact.get("email") or "").strip() or None

    # 学歴・職歴
    rows, history_overflow = build_history_rows(
        data.get("education") or [], data.get("work_history") or []
    )
    for (y_cell, m_cell, t_cell), row in zip(HISTORY_SLOTS, rows):
        y, m = _fmt_ym(row.get("year"), row.get("month"))
        ws[y_cell].value = y or None
        ws[m_cell].value = m or None
        _set_aligned(ws, t_cell, row.get("text"), row.get("align"))

    # 資格・免許
    quals = [q for q in (data.get("qualifications") or []) if (q.get("text") or "").strip()]
    qualification_overflow = max(0, len(quals) - MAX_QUALIFICATION_ROWS)
    for (y_cell, m_cell, t_cell), q in zip(QUALIFICATION_SLOTS, quals):
        y, m = _fmt_ym(q.get("year"), q.get("month"))
        ws[y_cell].value = y or None
        ws[m_cell].value = m or None
        ws[t_cell].value = q.get("text").strip()

    # 志望動機・本人希望
    ws[CELLS["motivation"]].value = (data.get("motivation") or "").strip() or None
    wishes = (data.get("wishes") or "").strip()
    if wishes:
        wish_lines = [ln for ln in wishes.splitlines() if ln.strip()] or [wishes]
        if len(wish_lines) > len(WISH_SLOTS):  # 行数超過はまとめて先頭ブロックへ
            wish_lines = ["\n".join(wish_lines)]
        for coord, line in zip(WISH_SLOTS, wish_lines):
            ws[coord].value = line
            ws[coord].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # 写真
    if photo_bytes:
        img = XLImage(_prepare_photo(photo_bytes))
        img.width, img.height = PHOTO_WIDTH_PX, PHOTO_HEIGHT_PX
        ws.add_image(img, PHOTO_ANCHOR)
    else:
        _draw_photo_placeholder(ws)

    # 印刷設定（A4 縦 2 枚・J 列で改ページはテンプレ設定を継承）
    ws.page_setup.scale = PRINT_SCALE
    ws.print_area = "B1:R86"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), {
        "history_overflow": history_overflow,
        "qualification_overflow": qualification_overflow,
    }


def soffice_available() -> bool:
    return shutil.which("soffice") is not None


def convert_xlsx_to_pdf(xlsx_bytes: bytes, timeout: int = 90) -> bytes | None:
    """LibreOffice で xlsx → PDF。soffice が無い/失敗したら None"""
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "rirekisho.xlsx"
            src.write_bytes(xlsx_bytes)
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=timeout,
            )
            pdf = Path(tmp) / "rirekisho.pdf"
            return pdf.read_bytes() if pdf.exists() else None
    except (subprocess.SubprocessError, OSError):
        return None
